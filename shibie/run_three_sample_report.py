from __future__ import annotations

"""Run the three representative samples and print a screenshot-ready report table.

This is a reporting wrapper only. All measurements still come from
``engine.pipeline.run_measurement`` and therefore use the existing main route.
"""

from pathlib import Path
import sys

from openpyxl import load_workbook

from engine.pipeline import run_measurement


PROJECT_ROOT = Path(__file__).resolve().parent
SAMPLE_NAMES = ("W088.jpg", "W114.jpg", "W054.jpg")


def fmt(value: object, digits: int = 3) -> str:
    if isinstance(value, (int, float)):
        return f"{value:.{digits}f}"
    return "-" if value is None else str(value)


def interval(values: object) -> str:
    if not isinstance(values, (list, tuple)) or len(values) != 2:
        return "-"
    return f"{fmt(values[0], 1)}-{fmt(values[1], 1)}"


def width_text(axis: dict) -> str:
    return "/".join(fmt(axis.get(key), 3) for key in ("p10_mm", "p50_mm", "p90_mm"))


def make_table(headers: list[str], rows: list[list[str]]) -> str:
    widths = [len(header) for header in headers]
    for row in rows:
        for index, cell in enumerate(row):
            widths[index] = max(widths[index], len(cell))

    def line(fill: str = "-") -> str:
        return "+" + "+".join(fill * (width + 2) for width in widths) + "+"

    def render(row: list[str]) -> str:
        return "|" + "|".join(f" {cell:<{width}} " for cell, width in zip(row, widths)) + "|"

    return "\n".join([line(), render(headers), line("="), *(render(row) for row in rows), line()])


def reference_labels() -> dict[str, tuple[object, object]]:
    workbook = load_workbook(
        PROJECT_ROOT / "tmp" / "fabricnet" / "FabricNet_parameters.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    labels = {
        f"W{int(row[0]):03d}.jpg": (row[2], row[3])
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    }
    workbook.close()
    return labels


def run_one(image_name: str, references: dict[str, tuple[object, object]]) -> list[str]:
    raw = (PROJECT_ROOT / "tmp" / "user_samples" / image_name).read_bytes()
    result = run_measurement(
        raw,
        raw,
        mm_per_px_x=0.0078125,
        mm_per_px_y=0.0083333,
        warp_axis_hint="vertical",
    )
    width = result.get("apparent_width") or {}
    twill = result.get("twill") or {}
    yarn = result.get("yarn_advisory") or {}
    gate3 = result.get("gate3") or {}
    warp_yarn = yarn.get("warp") or {}
    weft_yarn = yarn.get("weft") or {}
    ref_warp, ref_weft = references[image_name]

    weave = twill.get("weave_candidate") or "无明确候选"
    yarn_text = (
        f"经Ne {interval(warp_yarn.get('range_ne'))}/tex {interval(warp_yarn.get('range_tex'))}; "
        f"纬Ne {interval(weft_yarn.get('range_ne'))}/tex {interval(weft_yarn.get('range_tex'))}"
    )
    return [
        image_name,
        f"{ref_warp}/{ref_weft}",
        f"{fmt(result.get('warp_density_per_cm'))}/{fmt(result.get('weft_density_per_cm'))}",
        fmt(result.get("warp_weft_ratio")),
        result.get("status", "-"),
        f"{width_text(width.get('warp') or {})} ({width.get('status', '-')})",
        width_text(width.get("weft") or {}),
        f"{'左斜' if twill.get('direction') == 'left' else '右斜' if twill.get('direction') == 'right' else '-'} ({twill.get('direction_status', '-')})",
        f"{fmt(twill.get('angle_deg'))}° ({twill.get('angle_status', '-')})",
        f"{weave} ({twill.get('weave_status', '-')})",
        f"{yarn_text} [粗略参考]",
        gate3.get("status", "-"),
    ]


def main() -> None:
    # Ensure Chinese headers and statuses display correctly in the PyCharm terminal.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    references = reference_labels()
    rows = [run_one(name, references) for name in SAMPLE_NAMES]
    headers = [
        "样本",
        "参考经/纬",
        "实测经/纬(根/cm)",
        "经纬比",
        "密度状态",
        "经纱P10/P50/P90(mm)",
        "纬纱P10/P50/P90(mm)",
        "斜向",
        "角度",
        "组织候选",
        "Ne/tex粗略参考",
        "Gate3",
    ]
    print("三样本第六节测试结果（P10/P50/P90 顺序；同图双输入，mm/px=0.0078125/0.0083333）")
    print(make_table(headers, rows))
    print("说明：PASS=当前模块证据满足输出条件；LOW_CONFIDENCE=结果可输出，但保留质量提示；Ne/tex 为参考建议。")


if __name__ == "__main__":
    main()
