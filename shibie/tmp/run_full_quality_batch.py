from __future__ import annotations

import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook
from engine.pipeline import run_measurement


def compact_axis(axis: dict | None) -> dict:
    axis = axis or {}
    return {
        "status": axis.get("status"),
        "p10_mm": axis.get("p10_mm"), "p50_mm": axis.get("p50_mm"), "p90_mm": axis.get("p90_mm"),
        "valid_sample_count": axis.get("valid_sample_count", 0),
        "valid_yarn_count": axis.get("valid_yarn_count", 0),
        "reasons": axis.get("reasons", []),
    }


def main() -> None:
    sheet = load_workbook("tmp/fabricnet/FabricNet_parameters.xlsx", read_only=True, data_only=True).active
    labels = {int(row[0]): str(row[4]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    results = []
    for number in range(1, 131):
        name = f"W{number:03d}.jpg"
        raw = (Path("tmp/user_samples") / name).read_bytes()
        result = run_measurement(raw, raw, mm_per_px_x=0.0078125, mm_per_px_y=0.0083333, warp_axis_hint="vertical")
        twill = result.get("twill", {})
        width = result.get("apparent_width") or {}
        yarn = result.get("yarn_advisory") or {}
        results.append({
            "image": name, "fabricnet_texture": labels[number],
            "direction_status": twill.get("direction_status"), "direction": twill.get("direction"),
            "angle_status": twill.get("angle_status"), "angle_deg": twill.get("angle_deg"),
            "weave_status": twill.get("weave_status"), "weave_candidate": twill.get("weave_candidate"),
            "weave_result_kind": twill.get("weave_result_kind"),
            "candidate_confidence": twill.get("candidate_confidence"),
            "weave_reasons": twill.get("weave_reasons", []),
            "width_status": width.get("status"),
            "warp_width": compact_axis(width.get("warp")), "weft_width": compact_axis(width.get("weft")),
            "yarn_status": yarn.get("status"), "yarn_reference_only": yarn.get("reference_only"),
            "warp_ne": (yarn.get("warp") or {}).get("range_ne"),
            "weft_ne": (yarn.get("weft") or {}).get("range_ne"),
        })
    Path("tmp/full_quality_batch_001_130.json").write_text(json.dumps({"count": 130, "samples": results}, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
