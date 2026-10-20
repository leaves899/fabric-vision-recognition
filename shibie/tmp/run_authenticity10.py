from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.pipeline import run_measurement


NAMES = ["W088.jpg", "W004.jpg", "W063.jpg", "W012.jpg", "W014.jpg", "W001.jpg", "W003.jpg", "W100.jpg", "W096.jpg", "W050.jpg"]
MM_X = 0.0078125
MM_Y = 0.0083333


def fabricnet_values() -> dict[str, dict[str, object]]:
    from openpyxl import load_workbook

    sheet = load_workbook(ROOT / "tmp" / "fabricnet" / "FabricNet_parameters.xlsx", read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out: dict[str, dict[str, object]] = {}
    for row in rows[1:]:
        values = dict(zip(headers, row))
        image_id = values.get("Image id")
        if image_id is not None:
            try:
                key = f"W{int(image_id):03d}"
            except (TypeError, ValueError):
                key = Path(str(image_id).strip()).stem
            out[key] = values
    return out


def density(measurement: dict, key: str) -> float | None:
    value = measurement.get(key)
    return float(value) if isinstance(value, (int, float)) else None


def spacing(measurement: dict, direction: str) -> dict:
    item = measurement.get("density", {}).get(direction, {})
    return {
        "median_spacing_px": item.get("median_spacing_px"),
        "spacing_cv": item.get("spacing_cv"),
        "valid_centerline_count": item.get("valid_centerline_count", item.get("valid_line_count")),
    }


def main() -> None:
    refs = fabricnet_values()
    baseline_payload = json.loads((ROOT / "tmp" / "quality_selected_peak_suppressed_050.json").read_text(encoding="utf-8"))
    baseline = {item["image"]: item["measurement"] for item in baseline_payload["samples"]}
    rows = []
    for index, name in enumerate(NAMES, 1):
        raw = (ROOT / "tmp" / "user_samples" / name).read_bytes()
        measurement = run_measurement(raw, raw, mm_per_px_x=MM_X, mm_per_px_y=MM_Y, warp_axis_hint="vertical")
        ref = refs.get(Path(name).stem, {})
        old = baseline.get(name, {})
        rows.append({
            "image": name,
            "reference_warp": ref.get("Warp"),
            "reference_weft": ref.get("Weft"),
            "warp_density_per_cm": density(measurement, "warp_density_per_cm"),
            "weft_density_per_cm": density(measurement, "weft_density_per_cm"),
            "warp_density_before": density(old, "warp_density_per_cm"),
            "weft_density_before": density(old, "weft_density_per_cm"),
            "warp_delta_vs_reference": (density(measurement, "warp_density_per_cm") - float(ref["Warp"])) if density(measurement, "warp_density_per_cm") is not None and isinstance(ref.get("Warp"), (int, float)) else None,
            "weft_delta_vs_reference": (density(measurement, "weft_density_per_cm") - float(ref["Weft"])) if density(measurement, "weft_density_per_cm") is not None and isinstance(ref.get("Weft"), (int, float)) else None,
            "warp_weft_ratio": measurement.get("warp_weft_ratio"),
            "status": measurement.get("status"),
            "warp_spacing": spacing(measurement, "warp"),
            "weft_spacing": spacing(measurement, "weft"),
        })
        print(f"{index}/10 {name}", flush=True)
    payload = {"count": len(rows), "mm_per_px_x": MM_X, "mm_per_px_y": MM_Y, "warp_axis_hint": "vertical", "samples": rows}
    out_json = ROOT / "tmp" / "quality_selected_authenticity10_profile_verified.json"
    out_csv = ROOT / "tmp" / "quality_selected_authenticity10_profile_verified.csv"
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    with out_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=[key for key, value in rows[0].items() if not isinstance(value, dict)])
        writer.writeheader()
        writer.writerows({key: value for key, value in row.items() if not isinstance(value, dict)} for row in rows)
    print(out_json)
    print(out_csv)


if __name__ == "__main__":
    main()
