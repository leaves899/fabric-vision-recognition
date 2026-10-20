from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent


def references() -> dict[str, tuple[float | None, float | None]]:
    sheet = load_workbook(
        ROOT / "tmp" / "fabricnet" / "FabricNet_parameters.xlsx",
        read_only=True,
        data_only=True,
    ).active
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    image_i, warp_i, weft_i = headers.index("Image id"), headers.index("Warp"), headers.index("Weft")
    result = {}
    for row in rows:
        if row[image_i] is not None:
            result[f"W{int(row[image_i]):03d}"] = (row[warp_i], row[weft_i])
    return result


def main() -> None:
    payload = json.loads((ROOT / "tmp" / "quality_selected_authenticity50.json").read_text(encoding="utf-8"))
    refs = references()
    records = []
    for sample in payload["samples"]:
        measurement = sample["measurement"]
        ref_warp, ref_weft = refs[Path(sample["image"]).stem]
        for axis, reference in (("warp", ref_warp), ("weft", ref_weft)):
            periodicity = measurement["density_debug"][axis]["periodicity"]
            records.append({
                "image": sample["image"], "axis": axis, "reference": float(reference),
                "density": float(measurement[f"{axis}_density_per_cm"]),
                "periodicity_spacing_px": periodicity.get("spacing_px"),
                "centerline_spacing_px": periodicity.get("centerline_spacing_px"),
                "periodicity_difference_pct": periodicity.get("spacing_difference_pct"),
                "autocorrelation_peak": periodicity.get("autocorrelation_peak"),
                "periodicity_status": periodicity["status"],
            })

    over = [r for r in records if r["density"] > r["reference"]]
    conflict = [
        r for r in records
        if isinstance(r["periodicity_difference_pct"], float)
        and r["periodicity_difference_pct"] > 30.0
        and isinstance(r["autocorrelation_peak"], float)
        and r["autocorrelation_peak"] >= 0.18
    ]
    targets = [r for r in conflict if r["density"] > r["reference"]]
    summary = {
        "axes": len(records), "density_above_reference": len(over),
        "periodicity_centerline_conflict": len(conflict),
        "conflict_and_density_above_reference": len(targets),
        "targets": targets,
    }
    out = ROOT / "tmp" / "periodicity50_analysis.json"
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(out)


if __name__ == "__main__":
    main()
