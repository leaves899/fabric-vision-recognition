from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from engine.pipeline import run_measurement


def main(start: int, end: int) -> None:
    sheet = load_workbook(
        "tmp/fabricnet/FabricNet_parameters.xlsx", read_only=True, data_only=True,
    ).active
    labels = {int(row[0]): str(row[4]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    results: list[dict] = []
    summary: dict[str, Counter] = defaultdict(Counter)
    for number in range(start, end + 1):
        name = f"W{number:03d}.jpg"
        raw = (Path("tmp/user_samples") / name).read_bytes()
        twill = run_measurement(
            raw, raw,
            mm_per_px_x=0.0078125,
            mm_per_px_y=0.0083333,
            warp_axis_hint="vertical",
        )["twill"]
        label = labels[number]
        summary[label][f"direction_{twill.get('direction_status', 'UNKNOWN')}"] += 1
        summary[label][f"angle_{twill.get('angle_status', 'UNKNOWN')}"] += 1
        results.append({
            "image": name,
            "fabricnet_texture": label,
            "status": twill["status"],
            "direction_status": twill.get("direction_status"),
            "angle_status": twill.get("angle_status"),
            "candidate_direction": twill.get("candidate_direction"),
            "candidate_angle_deg": twill.get("candidate_angle_deg"),
            "direction_reasons": twill.get("direction_reasons", []),
            "angle_reasons": twill.get("angle_reasons", []),
            "weave_type": twill.get("weave_type"),
            "weave_candidate": twill.get("weave_candidate"),
            "weave_result_kind": twill.get("weave_result_kind"),
            "candidate_confidence": twill.get("candidate_confidence"),
            "weave_status": twill.get("weave_status"),
            "weave_reasons": twill.get("weave_reasons", []),
            "weave_evidence": twill.get("interlacing_topology", {}).get("weave_evidence", {}),
            "reasons": twill.get("reasons", []),
            "band_evidence": twill.get("band_evidence", []),
        })
        print(f"{number}/{end}", flush=True)
    output = {
        "range": [start, end],
        "summary": {label: dict(counts) for label, counts in sorted(summary.items())},
        "samples": results,
    }
    Path(f"tmp/twill_batch_{start:03d}_{end:03d}.json").write_text(
        json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8",
    )


if __name__ == "__main__":
    main(int(sys.argv[1]), int(sys.argv[2]))
